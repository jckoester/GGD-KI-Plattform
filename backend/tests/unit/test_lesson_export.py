"""Tests für lesson_export.py: export_markdown (golden), PDF/DOCX (smoke)."""

import pytest
from app.planning.lesson_export import ExportPhase, LessonExport, export_markdown, _slugify


def _make_export(**overrides) -> LessonExport:
    defaults = dict(
        titel="Einführung in die Fotosynthese",
        titel_slug="einfuehrung-in-die-fotosynthese",
        stundenziel="SuS können die Grundgleichung der Fotosynthese erklären.",
        gruppe="10b",
        gruppe_slug="10b",
        datum="2026-09-15",
        start_period=3,
        periods=2,
        verfuegbare_min=90,
        ue_titel="Stoffwechsel",
        phasen=[
            ExportPhase(
                name="Einstieg",
                dauer_min=15,
                beschreibung="Aktivierung Vorwissen",
                prio="kern",
                sozialform="Partnerarbeit",
                methode="Fishbowl",
                material=["Folie 1"],
            ),
            ExportPhase(
                name="Erarbeitung",
                dauer_min=30,
                beschreibung="",
                prio="kern",
                sozialform="",
                methode="",
                material=[],
            ),
            ExportPhase(
                name="Sicherung",
                dauer_min=15,
                beschreibung="",
                prio="uebung",
                sozialform="Gruppenarbeit",
                methode="Tafelbild",
                material=["AB 1", "AB 2"],
            ),
        ],
        refs=[
            {"typ": "ik", "code": "3.1.2", "titel": "Fotosynthese verstehen", "partiell": False},
            {"typ": "pk", "code": "K1", "titel": "Kommunikation", "partiell": True},
        ],
    )
    defaults.update(overrides)
    return LessonExport(**defaults)


# ── Slugify ───────────────────────────────────────────────────────────────────


def test_slugify_basic():
    assert _slugify("Einstieg") == "einstieg"


def test_slugify_umlauts():
    assert _slugify("Übungsaufgaben") == "uebungsaufgaben"
    assert _slugify("Käfig & Öl") == "kaefig-oel"


def test_slugify_special_chars():
    assert _slugify("Test: Aufgabe 1!") == "test-aufgabe-1"


def test_slugify_empty():
    assert _slugify("") == "stunde"


def test_slugify_long():
    s = "a" * 100
    assert len(_slugify(s)) <= 60


# ── Markdown-Golden-File ──────────────────────────────────────────────────────


EXPECTED_MD_FRONTMATTER = """\
---
titel: Einführung in die Fotosynthese
datum: 2026-09-15
gruppe: 10b
ue: Stoffwechsel
verfuegbar_min: 90
kompetenzen: [3.1.2, "K1[…]"]
---"""


def test_markdown_frontmatter():
    md = export_markdown(_make_export())
    for line in EXPECTED_MD_FRONTMATTER.splitlines():
        assert line in md, f"Erwartete Zeile nicht gefunden: {line!r}"


def test_markdown_frontmatter_ist_gueltiges_yaml():
    """Bis 0.17 stand `K1[…]` ohne Anführungszeichen in der Liste — eine eckige Klammer in
    einer Flow-Liste ist für YAML kaputt, und Obsidian las den ganzen Kopf nicht."""
    import yaml

    kopf = yaml.safe_load(export_markdown(_make_export()).split("---", 2)[1])
    assert kopf["kompetenzen"] == ["3.1.2", "K1[…]"]


def test_markdown_stundenziel():
    md = export_markdown(_make_export())
    assert "**Stundenziel:**" in md
    assert "SuS können die Grundgleichung" in md


def test_markdown_phase_headers():
    md = export_markdown(_make_export())
    assert "## Einstieg (15′ · Kern)" in md
    assert "## Erarbeitung (30′ · Kern)" in md
    assert "## Sicherung (15′ · Übung)" in md


def test_markdown_zeitbudget_no_ueberhang():
    md = export_markdown(_make_export())
    assert "60′ geplant / 90′ verfügbar" in md
    assert "Überhang" not in md


def test_markdown_zeitbudget_with_ueberhang():
    export = _make_export(verfuegbare_min=50)
    md = export_markdown(export)
    assert "Überhang" in md
    assert "+10′" in md


def test_markdown_methode_included():
    md = export_markdown(_make_export())
    assert "**Methode:** Fishbowl" in md
    assert "**Methode:** Tafelbild" in md


def test_markdown_sozialform_included():
    md = export_markdown(_make_export())
    assert "**Sozialform:** Partnerarbeit" in md
    assert "**Sozialform:** Gruppenarbeit" in md


def test_markdown_material_included():
    md = export_markdown(_make_export())
    assert "**Material:** Folie 1" in md
    assert "**Material:** AB 1" in md
    assert "**Material:** AB 2" in md


def test_markdown_empty_phase_no_methode():
    md = export_markdown(_make_export())
    # "Erarbeitung" hat keine Methode → kein Material-/Methode-Block für diese Phase
    lines = md.splitlines()
    erarbeitung_idx = next(i for i, l in enumerate(lines) if "Erarbeitung" in l)
    sicherung_idx = next(i for i, l in enumerate(lines) if "## Sicherung" in l)
    between = lines[erarbeitung_idx + 1 : sicherung_idx]
    assert not any("**Methode" in l or "**Material" in l for l in between)


def test_markdown_no_refs_no_kompetenzen_key():
    export = _make_export(refs=[])
    md = export_markdown(export)
    assert "kompetenzen:" not in md


def test_markdown_partiell_suffix():
    md = export_markdown(_make_export())
    assert "K1[…]" in md


# ── Begriffe der Stunde (0.17, A1 Schritt 4) ─────────────────────────────────

_BEGRIFFE = [
    {"typ": "concept", "code": None, "titel": "Benzol", "partiell": False, "content_type": "stoffsteckbrief"},
    {"typ": "concept", "code": None, "titel": "Mesomerie", "partiell": False},
]


def _gemischt():
    return _make_export(refs=_make_export().refs + _BEGRIFFE)


def _frontmatter(md: str) -> list[str]:
    return md.split("---", 2)[1].strip().splitlines()


def test_markdown_begriffe_getrennt_von_kompetenzen():
    zeilen = _frontmatter(export_markdown(_gemischt()))
    assert 'kompetenzen: [3.1.2, "K1[…]"]' in zeilen
    assert "fachbegriffe: [Benzol, Mesomerie]" in zeilen


def test_markdown_ohne_begriffe_wie_bisher():
    """Ein Entwurf ohne Begriffe exportiert wie in 0.16 — ohne `fachbegriffe:`."""
    md = export_markdown(_make_export())
    assert "fachbegriffe:" not in md
    for line in EXPECTED_MD_FRONTMATTER.splitlines():
        assert line in md


def test_markdown_nur_begriffe_kein_kompetenzen_schluessel():
    zeilen = _frontmatter(export_markdown(_make_export(refs=_BEGRIFFE)))
    assert not any(z.startswith("kompetenzen:") for z in zeilen)
    assert "fachbegriffe: [Benzol, Mesomerie]" in zeilen


def test_markdown_frontmatter_bleibt_yaml():
    """Titel mit Komma oder Doppelpunkt zerlegten die Liste; `code: null` stand als „None" da."""
    import yaml

    md = export_markdown(_make_export(refs=[
        {"typ": "ik", "code": None, "titel": "Säuren, Basen: Neutralisation", "partiell": False},
        {"typ": "concept", "code": None, "titel": "Säure-Base-Reaktion, Brønsted", "partiell": False},
    ]))
    kopf = yaml.safe_load(md.split("---", 2)[1])
    assert kopf["kompetenzen"] == ["Säuren, Basen: Neutralisation"]
    assert kopf["fachbegriffe"] == ["Säure-Base-Reaktion, Brønsted"]
    assert "None" not in md


def test_begriffe_mit_aliasen():
    """Nennt die Stunde einen Begriff unter Aliasen (Plan § 26), stehen sie im Export dabei."""
    import yaml

    se = {"typ": "concept", "code": None, "titel": "elektrophile Substitution", "partiell": False,
          "aliase": ["σ-Komplex", "Arenium-Ion"]}
    md = export_markdown(_make_export(refs=[se]))
    kopf = yaml.safe_load(md.split("---", 2)[1])
    assert kopf["fachbegriffe"] == ["elektrophile Substitution (σ-Komplex, Arenium-Ion)"]


# ── DOCX smoke ────────────────────────────────────────────────────────────────


def test_export_docx_returns_bytes():
    pytest.importorskip("docx")
    from app.planning.lesson_export import export_docx

    result = export_docx(_make_export())
    assert isinstance(result, bytes)
    assert len(result) > 1000  # non-trivial docx
    # DOCX magic bytes (PK zip header)
    assert result[:2] == b"PK"


def test_export_docx_empty_phasen():
    pytest.importorskip("docx")
    from app.planning.lesson_export import export_docx

    export = _make_export(phasen=[])
    result = export_docx(export)
    assert isinstance(result, bytes)
    assert result[:2] == b"PK"


def _docx_bloecke(export) -> list[str]:
    """Der Körper des DOCX in Reihenfolge: Absatztexte, Tabellen als „<tabelle>"."""
    import io

    from docx import Document

    from app.planning.lesson_export import export_docx

    koerper = Document(io.BytesIO(export_docx(export))).element.body
    return ["<tabelle>" if el.tag.endswith("}tbl") else "".join(t.text or "" for t in el.iter() if t.tag.endswith("}t"))
            for el in koerper if el.tag.endswith(("}p", "}tbl"))]


def test_docx_kompetenzen_und_begriffe_unter_der_tabelle():
    """Bis 0.18 fehlten beide im DOCX; MD und PDF führten sie. Lage wie im PDF: unter der Tabelle."""
    pytest.importorskip("docx")
    se = {"typ": "concept", "code": None, "titel": "elektrophile Substitution", "partiell": False,
          "aliase": ["σ-Komplex", "Arenium-Ion"]}
    bloecke = _docx_bloecke(_make_export(refs=_make_export().refs + [se]))
    tabelle = bloecke.index("<tabelle>")
    assert bloecke.index("Kompetenzen: 3.1.2, K1[…]") > tabelle
    assert bloecke.index("Fachbegriffe: elektrophile Substitution (σ-Komplex, Arenium-Ion)") > tabelle


def test_docx_ohne_refs_keine_absaetze():
    pytest.importorskip("docx")
    bloecke = _docx_bloecke(_make_export(refs=[]))
    assert not any(b.startswith(("Kompetenzen:", "Fachbegriffe:")) for b in bloecke)


def test_docx_nur_begriffe():
    pytest.importorskip("docx")
    bloecke = _docx_bloecke(_make_export(refs=_BEGRIFFE))
    assert "Fachbegriffe: Benzol, Mesomerie" in bloecke
    assert not any(b.startswith("Kompetenzen:") for b in bloecke)


def test_docx_ist_a4():
    """Die Vorlage von python-docx ist US-Letter; das PDF ist A4."""
    pytest.importorskip("docx")
    import io

    from docx import Document

    from app.planning.lesson_export import export_docx

    section = Document(io.BytesIO(export_docx(_make_export()))).sections[0]
    assert round(section.page_width.cm, 1) == 21.0
    assert round(section.page_height.cm, 1) == 29.7


def test_docx_spaltenbreiten_wie_im_pdf():
    """Bis 0.18 waren alle fünf Spalten gleich breit; die Phasenspalte trennte Wörter."""
    pytest.importorskip("docx")
    import io

    from docx import Document

    from app.planning.lesson_export import _SPALTEN_ANTEILE, export_docx

    doc = Document(io.BytesIO(export_docx(_make_export())))
    tabelle = doc.tables[0]
    gesamt = sum(c.width for c in tabelle.columns)
    for spalte, anteil in zip(tabelle.columns, _SPALTEN_ANTEILE):
        assert abs(spalte.width / gesamt * 100 - anteil) < 0.5
        assert all(zelle.width == spalte.width for zelle in spalte.cells)
    assert tabelle.columns[2].width > 3 * tabelle.columns[1].width  # Phase ≫ Prio


def test_pdf_prio_farbe_schlaegt_zeilenstreifen():
    """`tr:nth-child(even) td` ist spezifischer als eine einzelne Klasse — bis 0.18 stand das
    weiße Prio-Kürzel in jeder zweiten Zeile auf hellem Grund. Die Prio-Regeln brauchen
    mindestens zwei Klassen."""
    import re
    from pathlib import Path

    css = (Path(__file__).parents[2] / "app/planning/templates/lesson.html").read_text(encoding="utf-8")
    assert re.search(r"tr:nth-child\(even\)\s+td\s*\{", css), "Zeilenstreifen weg? Dann Test anpassen."
    for prio in ("kern", "uebung", "vertiefung"):
        regeln = re.findall(rf"^\s*([^{{}}\n]*\.prio-{prio}\b[^{{}}\n]*)\{{", css, re.M)
        assert regeln, prio
        assert all(sel.count(".") >= 2 for sel in regeln), regeln


# ── PDF smoke ─────────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_export_pdf_returns_bytes():
    pytest.importorskip("weasyprint")
    from app.planning.lesson_export import export_pdf

    result = await export_pdf(_make_export())
    assert isinstance(result, bytes)
    assert result[:4] == b"%PDF"


class TestFormelnImStundenPdf:
    """Kompetenztitel koennen Formeln tragen (`… die Zahl \\(\\pi\\) …`).

    Angezeigt wird sonst `r.code`; faellt die Anzeige auf den Titel zurueck, stuende
    ohne Prärendern die TeX-Quelle im PDF.
    """

    async def _pdf_html(self, export) -> str:
        """Erzeugt das PDF und gibt die HTML zurueck, die weasyprint bekommen haette."""
        from unittest.mock import AsyncMock, MagicMock, patch

        from app.planning.lesson_export import export_pdf

        gefangen = {}

        def _html(string=None, base_url=None):
            gefangen["html"] = string
            doc = MagicMock()
            doc.write_pdf = MagicMock(return_value=b"%PDF-fake")
            return doc

        with patch("weasyprint.HTML", new=_html), patch(
            "app.render.sidecar.render_math",
            new=AsyncMock(return_value='<svg class="mj"><path d="M0"/></svg>'),
        ):
            await export_pdf(export)
        return gefangen["html"]

    @pytest.mark.asyncio
    async def test_titel_mit_formel_wird_gerendert(self):
        html = await self._pdf_html(_make_export(refs=[
            # Ohne `code` faellt die Anzeige auf den Titel zurueck — der reale Fall.
            {"typ": "ik", "code": "", "titel": r"3.1.2(10) die Zahl \(\pi\) erklären",
             "partiell": False},
        ]))
        assert '<svg class="mj">' in html
        assert "\\(" not in html
        assert "erklären" in html

    @pytest.mark.asyncio
    async def test_code_hat_weiterhin_vorrang(self):
        html = await self._pdf_html(_make_export(refs=[
            {"typ": "ik", "code": "3.1.2", "titel": "Langer Titel", "partiell": False},
        ]))
        assert "3.1.2" in html
        assert "Langer Titel" not in html

    @pytest.mark.asyncio
    async def test_partiell_marker_bleibt(self):
        html = await self._pdf_html(_make_export(refs=[
            {"typ": "pk", "code": "K1", "titel": "Kommunikation", "partiell": True},
        ]))
        assert "[…]" in html

    @pytest.mark.asyncio
    async def test_titel_wird_escapet(self):
        """Das Template gibt mit `| safe` aus."""
        html = await self._pdf_html(_make_export(refs=[
            {"typ": "ik", "code": "", "titel": "<script>alert(1)</script>", "partiell": False},
        ]))
        assert "<script>alert(1)</script>" not in html
        assert "&lt;script&gt;" in html

    @pytest.mark.asyncio
    async def test_ohne_kompetenzen_kein_block(self):
        html = await self._pdf_html(_make_export(refs=[]))
        assert "Kompetenzen:" not in html
        assert "Fachbegriffe:" not in html

    @pytest.mark.asyncio
    async def test_begriffe_im_eigenen_block(self):
        html = await self._pdf_html(_gemischt())
        kompetenzen, begriffe = html.split("Fachbegriffe:", 1)
        assert "Benzol" not in kompetenzen.split("Kompetenzen:", 1)[1]
        assert "Benzol" in begriffe and "Mesomerie" in begriffe

    @pytest.mark.asyncio
    async def test_begriff_mit_aliasen(self):
        html = await self._pdf_html(_make_export(refs=[
            {"typ": "concept", "titel": "elektrophile Substitution", "aliase": ["σ-Komplex"]}]))
        assert "elektrophile Substitution (σ-Komplex)" in html.split("Fachbegriffe:", 1)[1]

    @pytest.mark.asyncio
    async def test_nur_begriffe_kein_kompetenzblock(self):
        html = await self._pdf_html(_make_export(refs=_BEGRIFFE))
        assert "Kompetenzen:" not in html
        assert "Fachbegriffe:" in html


# ── Phasen ohne Dauer (0.13, P1) ──────────────────────────────────────────────

def _mit_skizze(**overrides):
    """Die Beispielstunde, deren Erarbeitung noch keine Minuten hat."""
    export = _make_export(**overrides)
    export.phasen[1].dauer_min = None
    return export


def test_markdown_phase_ohne_dauer_nur_mit_prio():
    md = export_markdown(_mit_skizze())
    assert "## Erarbeitung (Kern)" in md
    assert "None" not in md and "0′ ·" not in md


def test_markdown_zeitbudget_zaehlt_nur_gesetzte_dauern_und_nennt_den_rest():
    md = export_markdown(_mit_skizze())
    # 15 + 15, die Erarbeitung zählt nicht als 0 still mit
    assert "30′ geplant / 90′ verfügbar · 1 Phase ohne Dauer" in md


def test_markdown_ohne_fehlende_dauer_kein_zusatz():
    assert "ohne Dauer" not in export_markdown(_make_export())


def test_docx_mit_phase_ohne_dauer():
    import io

    from docx import Document

    from app.planning.lesson_export import export_docx

    doc = Document(io.BytesIO(export_docx(_mit_skizze())))
    zeiten = [zeile.cells[0].text for zeile in doc.tables[0].rows[1:]]
    # Die Skizze beginnt bei 15. Danach steht die Uhr: Wann die Sicherung beginnt, weiß
    # niemand. Bis 03.10.2026 stand hier „15–30′" — die fehlende Dauer als 0 gezählt.
    assert zeiten == ["0–15′", "ab 15′", "–"]
    budget = next(p.text for p in doc.paragraphs if p.text.startswith("Zeitbudget"))
    assert "30′ / 90′ verfügbar · 1 Phase ohne Dauer" in budget


@pytest.mark.asyncio
async def test_pdf_mit_phase_ohne_dauer():
    """Bis 0.13.0 rechnete `lesson.html` selbst und brach mit `int + None` ab (0.13.1)."""
    pytest.importorskip("weasyprint")
    from app.planning.lesson_export import export_pdf

    assert (await export_pdf(_mit_skizze()))[:4] == b"%PDF"


@pytest.mark.asyncio
async def test_pdf_zeitspalte_und_budget_ohne_dauer():
    import re

    html = await TestFormelnImStundenPdf()._pdf_html(_mit_skizze())
    zellen = re.findall(r"<tr>\s*<td>([^<]*)</td>", html)
    # Dieselbe Regel wie im DOCX: nach der Skizze steht die Uhr.
    assert zellen == ["0–15′", "ab 15′", "–"]
    assert re.search(r"30′ geplant / 90′ verfügbar\s+· 1 Phase ohne Dauer", html)


# ── Gestrichene Phasen (0.19, Schritt 1) ──────────────────────────────────────
# Jan, 18.10.2026: Zeitspalte „–", die Uhr läuft ohne sie weiter, „offen" zählt mit.

def _mit_gestrichener():
    """Einstieg 15 · Erarbeitung 30 (gestrichen) · Sicherung 15."""
    export = _make_export()
    export.phasen[1].gestrichen = True
    return export


def test_zeitspalte_ueberspringt_gestrichene():
    from app.planning.lesson_export import zeitspalte

    # Die Sicherung beginnt, wo der Einstieg endet — bis 0.19 „45–60′".
    assert zeitspalte(_mit_gestrichener().phasen) == ["0–15′", "–", "15–30′"]


def test_zeitbudget_ohne_gestrichene():
    from app.planning.lesson_export import zeitbudget

    assert zeitbudget(_mit_gestrichener().phasen) == (30, 0)


def test_gestrichene_ohne_dauer_zaehlt_nicht_als_fehlend():
    from app.planning.lesson_export import zeitbudget

    export = _mit_gestrichener()
    export.phasen[1].dauer_min = None
    assert zeitbudget(export.phasen) == (30, 0)


def test_markdown_gestrichene_sichtbar_und_nicht_gezaehlt():
    md = export_markdown(_mit_gestrichener())
    assert "## ~~Erarbeitung~~ (30′ · Kern, gestrichen)" in md
    assert "30′ geplant / 90′ verfügbar" in md


def test_docx_gestrichene_durchgestrichen_ohne_zeit():
    import io

    from docx import Document

    from app.planning.lesson_export import export_docx

    doc = Document(io.BytesIO(export_docx(_mit_gestrichener())))
    zeilen = doc.tables[0].rows[1:]
    assert [z.cells[0].text for z in zeilen] == ["0–15′", "–", "15–30′"]
    name = zeilen[1].cells[2].paragraphs[0]
    assert name.text.startswith("Erarbeitung (gestrichen)")
    assert all(lauf.font.strike for lauf in name.runs)
    budget = next(p.text for p in doc.paragraphs if p.text.startswith("Zeitbudget"))
    assert "30′ / 90′ verfügbar" in budget


@pytest.mark.asyncio
async def test_pdf_gestrichene():
    import re

    html = await TestFormelnImStundenPdf()._pdf_html(_mit_gestrichener())
    assert re.findall(r"<tr>\s*<td>([^<]*)</td>", html) == ["0–15′", "–", "15–30′"]
    assert 'class="phase-name gestrichen"' in html
    assert re.search(r"30′ geplant / 90′ verfügbar", html)


def test_geplant_min_fuer_den_assistenten():
    from app.planning.phasen import geplant_min

    phasen = [{"dauer_min": 15}, {"dauer_min": 30, "status": "gestrichen"},
              {"dauer_min": 10, "status": "offen"}, {"dauer_min": None}, "kaputt"]
    assert geplant_min(phasen) == 25
    assert geplant_min(None) == 0
